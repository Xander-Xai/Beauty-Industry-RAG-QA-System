"""Multi-format DocumentProcessor tests using tiny local fixtures."""

from __future__ import annotations

import pytest

from offline.document_processor import DocumentProcessor
from offline.text_ingestion import DocumentProcessor as Phase1DocumentProcessor


def _txt(tmp_path, text="hello world policy content"):
    path = tmp_path / "doc.txt"
    path.write_text(text, encoding="utf-8")
    return path


def _pdf_text(tmp_path):
    import pymupdf

    path = tmp_path / "text.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Collagen moisturizer regulatory guidance for cosmetics industry.")
    document.save(path)
    document.close()
    return path


def _pdf_scanned(tmp_path):
    import pymupdf
    from PIL import Image

    image = Image.new("RGB", (200, 80), color="white")
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    path = tmp_path / "scanned.pdf"
    document = pymupdf.open()
    page = document.new_page(width=220, height=100)
    page.insert_image(pymupdf.Rect(10, 10, 210, 90), filename=str(image_path))
    document.save(path)
    document.close()
    return path


def _docx(tmp_path):
    from docx import Document

    path = tmp_path / "doc.docx"
    document = Document()
    document.add_heading("Safety Guidance", level=1)
    document.add_paragraph("First paragraph about ingredients.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ingredient"
    table.rows[0].cells[1].text = "Limit"
    document.save(path)
    return path


def _xlsx(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "doc.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ingredients"
    sheet.append(["Name", "Limit"])
    sheet.append(["Retinol", 0.3])
    sheet.append(["", None])
    sheet.append(["Niacinamide", 5])
    workbook.save(path)
    return path


def test_txt_processor_matches_phase_one_chunk_identity(tmp_path):
    path = _txt(tmp_path, "Alpha beta gamma delta " * 30)
    new = DocumentProcessor(chunk_size=50, chunk_overlap=5)
    old = Phase1DocumentProcessor(chunk_size=50, chunk_overlap=5)
    new_chunks = new.process_chunks(path, role_mask=0, dept_mask=0, doc_version_epoch="phase_1")
    old_chunks = old.process(path, role_mask=0, dept_mask=0, doc_version_epoch="phase_1")
    assert [chunk.chunk_id for chunk in new_chunks] == [chunk.chunk_id for chunk in old_chunks]
    assert [chunk.text for chunk in new_chunks] == [chunk.text for chunk in old_chunks]


def test_txt_bom_and_empty_document(tmp_path):
    path = tmp_path / "bom.txt"
    path.write_text("\ufeffsafety first", encoding="utf-8")
    processed = DocumentProcessor().process(path)
    assert processed.document_type == "txt"
    assert processed.blocks[0].text == "safety first"
    assert not processed.requires_ocr

    empty = tmp_path / "empty.txt"
    empty.write_text("  \n", encoding="utf-8")
    assert DocumentProcessor().process(empty).blocks == []


def test_pdf_text_extraction_records_page_metadata(tmp_path):
    processed = DocumentProcessor().process(_pdf_text(tmp_path))
    assert processed.document_type == "pdf"
    assert not processed.requires_ocr
    assert processed.blocks
    assert processed.blocks[0].metadata["page_number"] == 1
    assert "Collagen" in processed.blocks[0].text


def test_scanned_pdf_requests_ocr_instead_of_empty_success(tmp_path):
    processed = DocumentProcessor().process(_pdf_scanned(tmp_path))
    assert processed.requires_ocr is True
    assert processed.images
    assert processed.images[0].mime_type == "image/png"
    assert processed.images[0].page_number == 1


def test_docx_preserves_structure(tmp_path):
    processed = DocumentProcessor().process(_docx(tmp_path))
    block_types = [block.block_type for block in processed.blocks]
    assert block_types == ["heading", "paragraph", "table"]
    assert processed.blocks[0].metadata["heading_level"] == 1
    assert "Ingredient" in processed.blocks[-1].text


def test_xlsx_uses_sheet_row_windows(tmp_path):
    processor = DocumentProcessor(xlsx_rows_per_block=2)
    processed = processor.process(_xlsx(tmp_path))
    assert processed.document_type == "xlsx"
    assert processed.blocks
    first = processed.blocks[0]
    assert first.metadata["sheet_name"] == "Ingredients"
    assert first.metadata["row_start"] == 1
    assert first.metadata["row_end"] == 2
    assert "Retinol" in first.text


def test_unsupported_extension_and_limits(tmp_path):
    path = tmp_path / "doc.csv"
    path.write_text("a,b", encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported document type"):
        DocumentProcessor().process(path)

    oversized = tmp_path / "big.txt"
    oversized.write_text("x" * 20, encoding="utf-8")
    with pytest.raises(ValueError, match="ingestion limit"):
        DocumentProcessor(max_document_bytes=8).process(oversized)


@pytest.mark.parametrize("role,dept", [(-1, 0), (0, 2**32), (True, 0), (0, "1")])
def test_invalid_permission_masks_fail_closed(tmp_path, role, dept):
    path = _txt(tmp_path)
    with pytest.raises(ValueError, match="uint32"):
        DocumentProcessor().process_chunks(path, role_mask=role, dept_mask=dept, doc_version_epoch="epoch_1")


def test_chunk_limit_enforced_before_write(tmp_path):
    path = _txt(tmp_path, "abcdefghij" * 10)
    processor = DocumentProcessor(chunk_size=5, chunk_overlap=0, max_chunks=2)
    with pytest.raises(ValueError, match="2-chunk ingestion limit"):
        processor.process_chunks(path, role_mask=0, dept_mask=0, doc_version_epoch="epoch_1")


def test_source_identity_is_stable_across_roots(tmp_path):
    root_a = tmp_path / "a"
    root_b = tmp_path / "b"
    (root_a / "guides").mkdir(parents=True)
    (root_b / "guides").mkdir(parents=True)
    (root_a / "guides" / "p.txt").write_text("same", encoding="utf-8")
    (root_b / "guides" / "p.txt").write_text("same", encoding="utf-8")
    _, first = DocumentProcessor(source_root=root_a).document_identity(root_a / "guides" / "p.txt")
    _, second = DocumentProcessor(source_root=root_b).document_identity(root_b / "guides" / "p.txt")
    assert first == second
