"""Deterministic end-to-end offline pipeline coverage (no external services).

Covers a small multi-format corpus (TXT/PDF/DOCX/XLSX/image), dense + CLIP
retrieval, RBAC, epoch isolation, and failure handling where a failed staging
epoch must not be sealed.
"""

from __future__ import annotations

import io

import numpy as np
import pytest
from PIL import Image
from qdrant_client import QdrantClient

from offline.document_processor import DocumentProcessor
from offline.elasticsearch_writer import ElasticsearchWriter
from offline.embeddings import DeterministicTestEmbedder, DeterministicTestImageEmbedder
from offline.image_processor import DeterministicTestOCRProvider, ImageProcessor
from offline.qdrant_writer import QdrantImageWriter
from offline.snapshot_builder import IngestionSource, SnapshotBuilder
from offline.state_store import StateStore
from offline.text_ingestion import QdrantTextWriter
from tests.offline.fakes import FakeElasticsearchClient

DIMENSION = 8


def _png_bytes(color="white") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (48, 24), color=color).save(buffer, format="PNG")
    return buffer.getvalue()


def _pdf(tmp_path):
    import pymupdf

    path = tmp_path / "regulation.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Regulatory guidance for collagen cosmetics ingredients.")
    document.save(path)
    document.close()
    return path


def _docx(tmp_path):
    from docx import Document

    path = tmp_path / "formula.docx"
    document = Document()
    document.add_heading("Formula", level=1)
    document.add_paragraph("Niacinamide serum formulation steps.")
    document.save(path)
    return path


def _xlsx(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "limits.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Limits"
    sheet.append(["Ingredient", "Max"])
    sheet.append(["Retinol", 0.3])
    workbook.save(path)
    return path


def _builder(tmp_path, *, text_embedder=None, image_embedder=None, ocr_provider=None, es_writer=None):
    client = QdrantClient(":memory:")
    text_embedder = text_embedder or DeterministicTestEmbedder(dimension=DIMENSION)
    image_processor = ImageProcessor(
        ocr_provider or DeterministicTestOCRProvider(text="label retinol"),
        image_embedder or DeterministicTestImageEmbedder(dimension=DIMENSION),
        visual_weight_repeat=2,
    )
    if es_writer is None:
        es_writer = ElasticsearchWriter(FakeElasticsearchClient(), "cosmetics_docs")
        es_writer.ensure_index()
    builder = SnapshotBuilder(
        processor=DocumentProcessor(chunk_size=40, chunk_overlap=0, source_root=tmp_path),
        text_embedder=text_embedder,
        image_processor=image_processor,
        text_writer=QdrantTextWriter(client, "text_col", dimension=DIMENSION),
        image_writer=QdrantImageWriter(client, "image_col", dimension=DIMENSION),
        state_store=StateStore(tmp_path / "state.sqlite3"),
        es_writer=es_writer,
    )
    return builder, client


def _source(path, source_id, role=0, dept=0, document_type="text"):
    return IngestionSource(
        source_id=source_id,
        path=str(path),
        document_type=document_type,
        role_mask=role,
        dept_mask=dept,
        relative_path=source_id,
    )


def _corpus(tmp_path):
    txt = tmp_path / "public.txt"
    txt.write_text("Public collagen moisturizer guidance.", encoding="utf-8")
    image = tmp_path / "product.png"
    image.write_bytes(_png_bytes())
    return [
        _source(txt, "public.txt", role=0, dept=0),
        _source(_pdf(tmp_path), "regulation.pdf", role=2, dept=4),
        _source(_docx(tmp_path), "formula.docx", role=1, dept=1),
        _source(_xlsx(tmp_path), "limits.xlsx", role=2, dept=4),
        _source(image, "product.png", role=0, dept=0, document_type="image"),
    ]


def test_multiformat_epoch_dense_clip_rbac_and_isolation(tmp_path, monkeypatch):
    builder, client = _builder(tmp_path)
    result = builder.build_full(_corpus(tmp_path), "epoch_1", seal=True)
    assert result.documents_processed == 5
    assert result.validation.ok
    assert result.sealed
    assert builder.text_writer._is_epoch_sealed("epoch_1")
    assert builder.image_writer._is_epoch_sealed("epoch_1")

    from auth.bitmask_rbac import build_qdrant_filter, build_qdrant_image_filter
    from models.embedding_service import EmbeddingService

    reader = EmbeddingService.__new__(EmbeddingService)
    reader._qdrant_client = client
    query_vector = np.array(DeterministicTestEmbedder(DIMENSION).embed_texts(["collagen guidance"])[0])

    from common.models import RecallResult
    from retrieval.parallel_recall import ParallelRecallManager

    hits = reader.search_qdrant_text(
        query_vector, collection_name="text_col", top_k=10, qdrant_filter=build_qdrant_filter(0, 0, "epoch_1")
    )
    assert hits
    assert {hit["metadata"]["doc_version_epoch"] for hit in hits} == {"epoch_1"}

    candidates = [
        RecallResult(
            doc_id=hit["doc_id"],
            content=hit["content"],
            score=hit["score"],
            source="dense_bge",
            metadata=hit["metadata"],
        )
        for hit in hits
    ]
    manager = ParallelRecallManager()
    public_authorized = manager._apply_rbac_filter(candidates, 0, 0)
    assert all((candidate.metadata.get("role_mask") or 0) == 0 for candidate in public_authorized)
    restricted_authorized = manager._apply_rbac_filter(candidates, 2, 4)
    assert len(restricted_authorized) >= len(public_authorized)

    # CLIP image retrieval.
    import models.embedding_service as embedding_module

    saved = embedding_module.config
    try:
        embedding_module.config = {
            **saved,
            "embedding": {**saved["embedding"], "image_clip": {"collection": "image_col"}},
        }
        image_hits = reader.search_qdrant_image(
            np.array(DeterministicTestImageEmbedder(DIMENSION).embed_images([_png_bytes()])[0]),
            top_k=5,
            qdrant_filter=build_qdrant_image_filter(0, 0, "epoch_1"),
        )
    finally:
        embedding_module.config = saved
    assert image_hits
    assert any(hit["metadata"]["role_mask"] == 0 for hit in image_hits)

    # Epoch isolation: a different epoch is empty.
    empty = reader.search_qdrant_text(
        query_vector, collection_name="text_col", top_k=10, qdrant_filter=build_qdrant_filter(0, 0, "epoch_2")
    )
    assert empty == []


def test_incremental_epoch_is_a_complete_snapshot(tmp_path):
    builder, client = _builder(tmp_path)
    sources = _corpus(tmp_path)
    builder.build_full(sources, "epoch_1")

    (tmp_path / "public.txt").write_text("Updated public guidance for cosmetics.", encoding="utf-8")
    result = builder.build_incremental(sources, "epoch_1", "epoch_2")
    assert result.validation.ok

    def epoch_doc_ids(epoch):
        records, _ = client.scroll("text_col", limit=1000, with_payload=True)
        return {
            point.payload["doc_id"]
            for point in records
            if (point.payload or {}).get("doc_type") == "text" and point.payload.get("doc_version_epoch") == epoch
        }

    assert epoch_doc_ids("epoch_1") == epoch_doc_ids("epoch_2")
    # Unchanged documents keep their logical chunk ids across epochs.
    records, _ = client.scroll("text_col", limit=1000, with_payload=True)
    by_epoch = {}
    for point in records:
        if (point.payload or {}).get("doc_type") != "text":
            continue
        by_epoch.setdefault(point.payload["doc_version_epoch"], set()).add(point.payload["chunk_id"])
    assert by_epoch["epoch_1"] & by_epoch["epoch_2"]


def test_failures_leave_staging_epoch_unsealed(tmp_path):
    class FailingTextEmbedder(DeterministicTestEmbedder):
        def embed_texts(self, texts):
            raise RuntimeError("BGE failure")

    class FailingOCR(DeterministicTestOCRProvider):
        def extract(self, image_bytes, *, page_number=None, image_index=0):
            raise RuntimeError("OCR failure")

    class FailingImageEmbedder(DeterministicTestImageEmbedder):
        def embed_images(self, images):
            raise RuntimeError("CLIP failure")

    class FailingES(ElasticsearchWriter):
        def replace_document(self, *args, **kwargs):
            raise RuntimeError("ES failure")

    source = tmp_path / "doc.txt"
    source.write_text("content", encoding="utf-8")
    image = tmp_path / "product.png"
    image.write_bytes(_png_bytes())

    cases = [
        (
            "bge",
            _builder(tmp_path, text_embedder=FailingTextEmbedder(dimension=DIMENSION)),
            [_source(source, "doc.txt")],
        ),
        (
            "es",
            _builder(tmp_path, es_writer=FailingES(FakeElasticsearchClient(), "cosmetics_docs")),
            [_source(source, "doc.txt")],
        ),
    ]
    for name, (builder, _client), sources in cases:
        with pytest.raises(RuntimeError):
            builder.build_full(sources, f"failed_{name}", seal=True)
        assert not builder.text_writer._is_epoch_sealed(f"failed_{name}")

    builder, _ = _builder(tmp_path, ocr_provider=FailingOCR(), image_embedder=FailingImageEmbedder(dimension=DIMENSION))
    with pytest.raises(RuntimeError):
        builder.build_full([_source(image, "product.png", document_type="image")], "failed_ocr", seal=True)
    assert not builder.image_writer._is_epoch_sealed("failed_ocr")
