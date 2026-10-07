"""Multi-format offline document processor.

Supported source formats: UTF-8 ``.txt``, ``.pdf``, ``.docx`` and ``.xlsx``.
Text and structure are extracted into ordered :class:`DocumentBlock` objects;
pages that contain little or no extractable text are exposed as
:class:`ExtractedImage` render requests so the OCR/image pipeline can process
them instead of silently returning empty content.

Heavy parsing libraries (PyMuPDF, python-docx, openpyxl) are imported lazily so
that importing this module never loads native or model dependencies.
"""

from __future__ import annotations

import hashlib
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from offline.chunking import POINT_NAMESPACE, chunk_blocks, chunk_identity
from offline.validation import (
    ensure_within_root,
    validate_epoch,
    validate_permissions,
    validate_source_id,
)

SUPPORTED_DOCUMENT_EXTENSIONS = (".txt", ".pdf", ".docx", ".xlsx")
_DOCUMENT_TYPE_BY_EXTENSION = {
    ".txt": "txt",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
}


@dataclass(frozen=True)
class DocumentBlock:
    """One ordered structural unit extracted from a document."""

    block_type: str
    text: str
    order: int
    block_identity: str
    metadata: dict = field(default_factory=dict)


@dataclass(frozen=True)
class ExtractedImage:
    """An image (or rendered scanned page) that must go through OCR/CLIP."""

    image_id: str
    data: bytes
    mime_type: str
    page_number: int | None
    image_index: int
    metadata: dict = field(default_factory=dict)


@dataclass(frozen=True)
class ProcessedDocument:
    """Parsed document content before chunking or embedding."""

    doc_id: str
    source_id: str
    source_path: str
    document_type: str
    content_hash: str
    blocks: list[DocumentBlock]
    images: list[ExtractedImage]
    requires_ocr: bool
    metadata: dict = field(default_factory=dict)


@dataclass(frozen=True)
class TextChunk:
    """A bounded text unit plus its storage and authorization contract."""

    doc_id: str
    chunk_id: str
    source_path: str
    text: str
    chunk_index: int
    content_hash: str
    role_mask: int
    dept_mask: int
    status: str
    doc_version_epoch: str
    metadata: dict = field(default_factory=dict)
    #: Canonical ingestion provenance (``offline.source_trust.SourceTrustRecord``),
    #: persisted verbatim onto the point. Empty means "not declared", which the
    #: writer refuses rather than passing through as managed content.
    provenance: dict = field(default_factory=dict)


def _normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


class DocumentProcessor:
    """Parse TXT/PDF/DOCX/XLSX into ordered blocks and image requests."""

    def __init__(
        self,
        chunk_size: int = 500,
        chunk_overlap: int = 50,
        max_document_bytes: int = 131_072,
        max_chunks: int = 256,
        source_root: str | Path | None = None,
        max_binary_document_bytes: int = 67_108_864,
        ocr_min_text_chars: int = 20,
        pdf_render_dpi: int = 150,
        max_images_per_document: int = 100,
        xlsx_rows_per_block: int = 50,
    ):
        if type(chunk_size) is not int or chunk_size <= 0:
            raise ValueError("chunk_size must be a positive integer")
        if type(chunk_overlap) is not int or not 0 <= chunk_overlap < chunk_size:
            raise ValueError("chunk_overlap must be an integer in [0, chunk_size)")
        if type(max_document_bytes) is not int or max_document_bytes <= 0:
            raise ValueError("max_document_bytes must be a positive integer")
        if type(max_chunks) is not int or max_chunks <= 0:
            raise ValueError("max_chunks must be a positive integer")
        for name, value in (
            ("max_binary_document_bytes", max_binary_document_bytes),
            ("ocr_min_text_chars", ocr_min_text_chars),
            ("pdf_render_dpi", pdf_render_dpi),
            ("max_images_per_document", max_images_per_document),
            ("xlsx_rows_per_block", xlsx_rows_per_block),
        ):
            if type(value) is not int or value <= 0:
                raise ValueError(f"{name} must be a positive integer")
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.max_document_bytes = max_document_bytes
        self.max_chunks = max_chunks
        self.source_root = Path(source_root).resolve() if source_root is not None else None
        self.max_binary_document_bytes = max_binary_document_bytes
        self.ocr_min_text_chars = ocr_min_text_chars
        self.pdf_render_dpi = pdf_render_dpi
        self.max_images_per_document = max_images_per_document
        self.xlsx_rows_per_block = xlsx_rows_per_block

    # -- identity ---------------------------------------------------------

    def document_identity(self, source: str | Path, source_id: str | None = None) -> tuple[str, str]:
        resolved_path = Path(source).resolve()
        source_path = str(resolved_path)
        if source_id is not None:
            identity = validate_source_id(source_id)
        elif self.source_root is not None:
            identity = ensure_within_root(resolved_path, self.source_root)
        else:
            identity = source_path
        doc_id = hashlib.sha256(identity.encode("utf-8")).hexdigest()
        return source_path, doc_id

    @staticmethod
    def detect_document_type(source: str | Path) -> str:
        suffix = Path(source).suffix.lower()
        document_type = _DOCUMENT_TYPE_BY_EXTENSION.get(suffix)
        if document_type is None:
            supported = ", ".join(SUPPORTED_DOCUMENT_EXTENSIONS)
            raise ValueError(f"unsupported document type {suffix or '<none>'}; supported: {supported}")
        return document_type

    # -- parsing ----------------------------------------------------------

    def process(self, source: str | Path, *, source_id: str | None = None) -> ProcessedDocument:
        path = Path(source)
        document_type = self.detect_document_type(path)
        limit = self.max_document_bytes if document_type == "txt" else self.max_binary_document_bytes
        raw = self._read_bytes(path, limit)
        source_path, doc_id = self.document_identity(path, source_id)
        resolved_source_id = source_id if source_id is not None else source_path
        if document_type == "txt":
            blocks, images, content_hash, requires_ocr = self._process_txt(raw, doc_id, source_path)
        elif document_type == "pdf":
            blocks, images, content_hash, requires_ocr = self._process_pdf(path, raw, doc_id, source_path)
        elif document_type == "docx":
            blocks, images, content_hash, requires_ocr = self._process_docx(path, raw)
        else:
            blocks, images, content_hash, requires_ocr = self._process_xlsx(path, raw)
        return ProcessedDocument(
            doc_id=doc_id,
            source_id=resolved_source_id,
            source_path=source_path,
            document_type=document_type,
            content_hash=content_hash,
            blocks=blocks,
            images=images,
            requires_ocr=requires_ocr,
            metadata={"source_path": source_path, "format": document_type},
        )

    def _read_bytes(self, path: Path, limit: int) -> bytes:
        try:
            with path.open("rb") as handle:
                raw = handle.read(limit + 1)
        except OSError as exc:
            raise ValueError(f"cannot read document {path}: {exc}") from exc
        if len(raw) > limit:
            raise ValueError(f"document exceeds the {limit}-byte ingestion limit: {path}")
        return raw

    def _process_txt(
        self, raw: bytes, doc_id: str, source_path: str
    ) -> tuple[list[DocumentBlock], list[ExtractedImage], str, bool]:
        try:
            text = _normalize_text(raw.decode("utf-8-sig"))
        except UnicodeError as exc:
            raise ValueError(f"cannot read UTF-8 text document {source_path}: {exc}") from exc
        if not text:
            return [], [], hashlib.sha256(b"").hexdigest(), False
        content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
        block = DocumentBlock(
            block_type="text",
            text=text,
            order=0,
            block_identity="",
            metadata={"format": "txt"},
        )
        return [block], [], content_hash, False

    def _process_pdf(
        self, path: Path, raw: bytes, doc_id: str, source_path: str
    ) -> tuple[list[DocumentBlock], list[ExtractedImage], str, bool]:
        pymupdf = _import_pymupdf()
        blocks: list[DocumentBlock] = []
        images: list[ExtractedImage] = []
        requires_ocr = False
        try:
            document = pymupdf.open(path)
        except Exception as exc:  # pragma: no cover - library-specific error type
            raise ValueError(f"cannot open PDF document {path}: {exc}") from exc
        try:
            for page_index, page in enumerate(document):
                page_text = _normalize_text(page.get_text("text") or "")
                if len(page_text) >= self.ocr_min_text_chars:
                    blocks.append(
                        DocumentBlock(
                            block_type="page",
                            text=page_text,
                            order=page_index,
                            block_identity=f"page:{page_index}",
                            metadata={"page_number": page_index + 1, "format": "pdf"},
                        )
                    )
                    continue
                requires_ocr = True
                if len(images) >= self.max_images_per_document:
                    continue
                image = self._render_pdf_page(page, page_index, doc_id)
                if image is not None:
                    images.append(image)
        finally:
            document.close()
        content_hash = hashlib.sha256(raw).hexdigest()
        return blocks, images, content_hash, requires_ocr

    def _render_pdf_page(self, page, page_index: int, doc_id: str):
        try:
            pixmap = page.get_pixmap(dpi=self.pdf_render_dpi)
            data = pixmap.tobytes("png")
        except Exception:  # pragma: no cover - renderer specific failures
            return None
        image_id = str(uuid.uuid5(POINT_NAMESPACE, f"{doc_id}:page-render:{page_index}"))
        return ExtractedImage(
            image_id=image_id,
            data=data,
            mime_type="image/png",
            page_number=page_index + 1,
            image_index=page_index,
            metadata={"format": "pdf-page", "page_number": page_index + 1},
        )

    def _process_docx(self, path: Path, raw: bytes) -> tuple[list[DocumentBlock], list[ExtractedImage], str, bool]:
        from docx import Document

        try:
            document = Document(str(path))
        except Exception as exc:  # pragma: no cover - library-specific error type
            raise ValueError(f"cannot open DOCX document {path}: {exc}") from exc
        blocks: list[DocumentBlock] = []
        order = 0
        for item in _iter_docx_items(document):
            block = self._docx_item_to_block(item, order)
            if block is not None:
                blocks.append(block)
                order += 1
        content_hash = hashlib.sha256(raw).hexdigest()
        return blocks, [], content_hash, False

    def _docx_item_to_block(self, item, order: int):
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                return None
            style_name = (item.style.name if item.style is not None else "") or ""
            heading_level = _heading_level(style_name)
            if heading_level is not None:
                return DocumentBlock(
                    block_type="heading",
                    text=text,
                    order=order,
                    block_identity=f"heading:{order}",
                    metadata={"heading_level": heading_level, "paragraph_index": order, "format": "docx"},
                )
            return DocumentBlock(
                block_type="paragraph",
                text=text,
                order=order,
                block_identity=f"paragraph:{order}",
                metadata={"paragraph_index": order, "format": "docx"},
            )
        if isinstance(item, Table):
            rows = []
            for row in item.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = _normalize_text("\n".join(rows))
            if not text:
                return None
            return DocumentBlock(
                block_type="table",
                text=text,
                order=order,
                block_identity=f"table:{order}",
                metadata={"table_index": order, "format": "docx"},
            )
        return None

    def _process_xlsx(self, path: Path, raw: bytes) -> tuple[list[DocumentBlock], list[ExtractedImage], str, bool]:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        except Exception as exc:  # pragma: no cover - library-specific error type
            raise ValueError(f"cannot open XLSX document {path}: {exc}") from exc
        blocks: list[DocumentBlock] = []
        order = 0
        try:
            for sheet_index, sheet_name in enumerate(workbook.sheetnames):
                sheet = workbook[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                for start in range(0, len(rows), self.xlsx_rows_per_block):
                    window = rows[start : start + self.xlsx_rows_per_block]
                    lines = []
                    for row in window:
                        cells = [_cell_to_text(value) for value in row]
                        if any(cell for cell in cells):
                            lines.append(" | ".join(cells))
                    text = _normalize_text("\n".join(lines))
                    if not text:
                        continue
                    row_start = start + 1
                    row_end = start + len(window)
                    blocks.append(
                        DocumentBlock(
                            block_type="row_window",
                            text=text,
                            order=order,
                            block_identity=f"sheet:{sheet_index}:rows:{row_start}-{row_end}",
                            metadata={
                                "sheet_name": sheet_name,
                                "row_start": row_start,
                                "row_end": row_end,
                                "format": "xlsx",
                            },
                        )
                    )
                    order += 1
        finally:
            workbook.close()
        content_hash = hashlib.sha256(raw).hexdigest()
        return blocks, [], content_hash, False

    # -- chunking ---------------------------------------------------------

    def build_chunks(
        self,
        processed: ProcessedDocument,
        *,
        role_mask: int,
        dept_mask: int,
        doc_version_epoch: str,
        provenance: dict | None = None,
    ) -> list[TextChunk]:
        validate_permissions(role_mask, dept_mask)
        validate_epoch(doc_version_epoch)
        units = chunk_blocks(
            processed.blocks,
            doc_id=processed.doc_id,
            content_hash=processed.content_hash,
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
        )
        if len(units) > self.max_chunks:
            raise ValueError(f"document exceeds the {self.max_chunks}-chunk ingestion limit: {processed.source_path}")
        chunks = []
        for unit in units:
            chunks.append(
                TextChunk(
                    doc_id=processed.doc_id,
                    chunk_id=chunk_identity(
                        processed.doc_id,
                        processed.content_hash,
                        unit.block_identity,
                        unit.chunk_index,
                        unit.text,
                    ),
                    source_path=processed.source_path,
                    text=unit.text,
                    chunk_index=unit.chunk_index,
                    content_hash=processed.content_hash,
                    role_mask=role_mask,
                    dept_mask=dept_mask,
                    status="active",
                    doc_version_epoch=doc_version_epoch,
                    metadata={
                        **processed.metadata,
                        **unit.metadata,
                    },
                    provenance=dict(provenance or {}),
                )
            )
        return chunks

    def process_chunks(
        self,
        source: str | Path,
        *,
        role_mask: int,
        dept_mask: int,
        doc_version_epoch: str,
        source_id: str | None = None,
        provenance: dict | None = None,
    ) -> list[TextChunk]:
        validate_permissions(role_mask, dept_mask)
        validate_epoch(doc_version_epoch)
        processed = self.process(source, source_id=source_id)
        return self.build_chunks(
            processed,
            role_mask=role_mask,
            dept_mask=dept_mask,
            doc_version_epoch=doc_version_epoch,
            provenance=provenance,
        )


def _import_pymupdf():
    try:
        import pymupdf  # type: ignore

        return pymupdf
    except ImportError:
        import fitz  # type: ignore

        return fitz


def _iter_docx_items(document):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _heading_level(style_name: str) -> int | None:
    normalized = style_name.strip().lower()
    for prefix in ("heading ", "heading", "标题 ", "标题"):
        if normalized.startswith(prefix.strip()):
            remainder = normalized[len(prefix.strip()) :].strip()
            if remainder.isdigit():
                return int(remainder)
    return None


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
